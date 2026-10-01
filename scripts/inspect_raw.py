import openpyxl
import os
import glob

SAMPLE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "sample_data")

def inspect_raw():
    files = sorted(glob.glob(os.path.join(SAMPLE_DIR, "*.xlsx")))
    out_file = os.path.join(SAMPLE_DIR, "raw_inspection.txt")
    with open(out_file, "w", encoding="utf-8") as f_out:
        for f in files:
            wb_name = os.path.basename(f)
            f_out.write(f"\n==========================================\n")
            f_out.write(f"FILE: {wb_name}\n")
            f_out.write(f"==========================================\n")
            wb = openpyxl.load_workbook(f, data_only=True)
            for sname in wb.sheetnames:
                ws = wb[sname]
                f_out.write(f"\n--- SHEET: '{sname}' (max_row={ws.max_row}, max_col={ws.max_column}) ---\n")
                merged = [str(m) for m in ws.merged_cells.ranges]
                if merged:
                    f_out.write(f"  Merged cell ranges: {merged}\n")
                for r in range(1, min(ws.max_row + 1, 15)):
                    row_vals = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column + 1, 15))]
                    if any(v is not None for v in row_vals):
                        f_out.write(f"  Row {r:2d}: {row_vals}\n")

if __name__ == "__main__":
    inspect_raw()
