import openpyxl
import os
import glob
import re
import json

SAMPLE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "sample_data")

def deep_analyze():
    files = sorted(glob.glob(os.path.join(SAMPLE_DIR, "*.xlsx")))
    
    analysis_results = {}
    
    for fpath in files:
        fname = os.path.basename(fpath)
        wb = openpyxl.load_workbook(fpath, data_only=True)
        wb_data = {
            "sheets": {}
        }
        
        for sname in wb.sheetnames:
            ws = wb[sname]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
                
            # Header detection
            header_idx = -1
            for r_i, r_val in enumerate(rows[:10]):
                str_vals = [str(v).strip().lower() for v in r_val if v is not None]
                if any("reg" in s or "roll" in s or "name" in s or "s.no" in s or "s no" in s for s in str_vals):
                    header_idx = r_i
                    break
                    
            if header_idx == -1:
                header_idx = 0
                
            headers = [str(v).strip() if v is not None else "" for v in rows[header_idx]]
            
            # Find reg_no and name col
            reg_col = -1
            name_col = -1
            dept_col = -1
            
            for c_i, h in enumerate(headers):
                h_lower = h.lower()
                if any(k in h_lower for k in ["reg", "roll", "htno"]):
                    reg_col = c_i
                elif "name" in h_lower and "reg" not in h_lower:
                    name_col = c_i
                elif any(k in h_lower for k in ["dept", "department", "branch", "specialization"]):
                    dept_col = c_i
                    
            # Fallback for sheets like Soliton Round 4 / OFFER where Reg No is stored under S.No column or col 0/6
            if reg_col == -1:
                # Check data rows for 12-char reg_no pattern
                for c_i in range(len(headers)):
                    col_vals = [str(rows[r][c_i]).strip() for r in range(header_idx + 1, min(len(rows), header_idx + 10)) if c_i < len(rows[r]) and rows[r][c_i] is not None]
                    if any(re.match(r'^7376\d{2}[A-Z0-9]{5,6}$', v) for v in col_vals):
                        reg_col = c_i
                        break

            records = []
            duplicate_regs = []
            malformed_regs = []
            seen_regs = set()
            
            for r_i in range(header_idx + 1, len(rows)):
                r_val = rows[r_i]
                if not any(v is not None and str(v).strip() for v in r_val):
                    continue
                    
                reg_v = str(r_val[reg_col]).strip() if reg_col != -1 and reg_col < len(r_val) and r_val[reg_col] is not None else ""
                name_v = str(r_val[name_col]).strip() if name_col != -1 and name_col < len(r_val) and r_val[name_col] is not None else ""
                dept_v = str(r_val[dept_col]).strip() if dept_col != -1 and dept_col < len(r_val) and r_val[dept_col] is not None else ""
                
                # Check float format for reg_v
                if reg_v and isinstance(r_val[reg_col], float):
                    reg_v = str(int(r_val[reg_col]))
                    
                if reg_v:
                    if reg_v in seen_regs:
                        duplicate_regs.append((r_i + 1, reg_v, name_v))
                    else:
                        seen_regs.add(reg_v)
                        
                    # Pattern check for BITSATHY reg numbers: 7376 + 2 digits (year e.g. 23/24) + 2 letters (AD/CS/EC/IT/AL/CB/CD/SE/AG/ME) + digits/letters
                    if not re.match(r'^7376\d{2}[A-Za-z0-9]{5,6}$', reg_v):
                        malformed_regs.append((r_i + 1, reg_v, name_v))
                        
                records.append({
                    "row": r_i + 1,
                    "reg_no": reg_v,
                    "name": name_v,
                    "dept": dept_v
                })
                
            wb_data["sheets"][sname] = {
                "header_row_1based": header_idx + 1,
                "headers": headers,
                "reg_col_header": headers[reg_col] if reg_col != -1 and reg_col < len(headers) else None,
                "name_col_header": headers[name_col] if name_col != -1 and name_col < len(headers) else None,
                "dept_col_header": headers[dept_col] if dept_col != -1 and dept_col < len(headers) else None,
                "record_count": len(records),
                "unique_reg_count": len(seen_regs),
                "duplicate_regs": duplicate_regs,
                "malformed_regs": malformed_regs,
                "reg_numbers_list": list(seen_regs)
            }
            
        analysis_results[fname] = wb_data
        
    out_path = os.path.join(SAMPLE_DIR, "deep_analysis.json")
    with open(out_path, "w", encoding="utf-8") as out:
        json.dump(analysis_results, out, indent=2)
    print(f"Deep analysis saved to {out_path}")

if __name__ == "__main__":
    deep_analyze()
