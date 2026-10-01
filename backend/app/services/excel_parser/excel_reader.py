import openpyxl
from typing import Any
from pathlib import Path


def read_workbook_sheets_raw(file_path: str) -> dict[str, list[list[Any]]]:
    """
    Reads an Excel workbook file and returns a dictionary of sheet_name -> list of raw cell rows.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found at: {file_path}")

    wb = openpyxl.load_workbook(file_path, data_only=True)
    raw_data: dict[str, list[list[Any]]] = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            sheet_rows.append(list(row))
        raw_data[sheet_name] = sheet_rows

    wb.close()
    return raw_data
